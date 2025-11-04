"""
Excel处理工具
"""
import time
import random
from datetime import datetime
from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


def generate_order_no() -> str:
    """生成19位订单编号"""
    timestamp = int(time.time() * 1000)
    random_num = random.randint(1000, 9999)
    order_no = f"{timestamp}{random_num}"

    if len(order_no) > 19:
        order_no = order_no[:19]
    elif len(order_no) < 19:
        order_no = order_no + "0" * (19 - len(order_no))

    return order_no


def parse_excel_orders(file_content: bytes):
    """解析Excel文件中的订单数据（支持规格和供应商字段，支持优化版格式）"""
    # 尝试读取Excel文件，自动检测工作表
    try:
        # 先尝试读取所有工作表名称
        xl = pd.ExcelFile(BytesIO(file_content))
        sheet_names = xl.sheet_names

        # 优先使用CAISHENDAO工作表，否则使用第一个工作表
        sheet_name = "CAISHENDAO" if "CAISHENDAO" in sheet_names else sheet_names[0]

        df = pd.read_excel(BytesIO(file_content), sheet_name=sheet_name, header=None)

        # 检测数据格式
        # 如果第一行包含"日期"、"产品名称"等，说明是新格式（表头在第一行，数据从第二行开始）
        first_row = df.iloc[0].astype(str).tolist()

        if "日期" in first_row and "产品名称" in first_row:
            # 检查是否为优化版格式（有"付款状态"字段）
            if "付款状态" in first_row:
                # 优化版格式：日期、付款状态、订单编号、产品名称、规格、采购金额、供应商、时间
                df_data = df.iloc[1:, :8].copy()
                df_data.columns = ["日期", "付款状态", "订单编号", "产品名称", "规格", "采购金额", "供应商", "时间"]
                # 优化版标记
                df_data["_format"] = "optimized"
            # 检查是否包含规格和供应商字段（新格式v2）
            elif "规格" in first_row and "供应商" in first_row:
                # 新格式v2：包含规格和供应商，9列
                df_data = df.iloc[1:, :9].copy()
                df_data.columns = ["日期", "初始状态", "订单编号", "产品名称", "规格", "采购金额", "供应商", "时间", "先采后付"]
                df_data["_format"] = "v2"
            else:
                # 新格式v1：没有规格和供应商，7列
                df_data = df.iloc[1:, :7].copy()
                df_data.columns = ["日期", "初始状态", "订单编号", "产品名称", "采购金额", "时间", "先采后付"]
                # 添加空的规格和供应商列
                df_data["规格"] = None
                df_data["供应商"] = None
                df_data["_format"] = "v1"
        else:
            # 旧格式：数据在列10-16，从第7行开始
            df_data = df.iloc[7:, 10:17].copy()
            df_data.columns = ["日期", "初始状态", "订单编号", "产品名称", "采购金额", "时间", "先采后付"]
            # 添加空的规格和供应商列
            df_data["规格"] = None
            df_data["供应商"] = None
            df_data["_format"] = "old"

    except Exception as e:
        raise ValueError(f"无法读取Excel文件: {str(e)}")

    orders = []
    errors = []

    # 获取数据格式
    is_optimized = df_data.iloc[0]["_format"] == "optimized" if len(df_data) > 0 else False

    for idx, row in df_data.iterrows():
        try:
            # 验证必填字段
            if pd.isna(row["产品名称"]) or not str(row["产品名称"]).strip():
                errors.append({"row": idx, "error": "产品名称缺失"})
                continue

            if pd.isna(row["采购金额"]):
                errors.append({"row": idx, "error": "采购金额缺失"})
                continue

            if pd.isna(row["日期"]):
                errors.append({"row": idx, "error": "订单日期缺失"})
                continue

            # 处理订单编号
            order_no_value = row.get("订单编号")
            if pd.notna(order_no_value) and str(order_no_value).strip():
                try:
                    order_no = str(int(order_no_value)).strip()
                except:
                    order_no = str(order_no_value).strip()
            else:
                order_no = generate_order_no()

            # 处理日期
            order_date = pd.to_datetime(row["日期"]).date()
            record_time = pd.to_datetime(row["时间"]) if pd.notna(row["时间"]) else datetime.now()

            # 根据格式处理支付状态
            if is_optimized:
                # 优化版：使用"付款状态"字段
                payment_status = str(row["付款状态"]).strip() if pd.notna(row["付款状态"]) else "即时付款"

                # 向后兼容：为旧字段设置值
                if payment_status == "即时付款":
                    order_status = "已付款"
                    payment_method = "已付款"
                elif payment_status in ["账期未到", "账期已结"]:
                    order_status = "先采后付"
                    payment_method = "先采后付"
                else:
                    # 未知状态，默认为即时付款
                    payment_status = "即时付款"
                    order_status = "已付款"
                    payment_method = "已付款"

                # 对于账期订单，"日期"字段是确认收货日期
                receive_date = order_date if payment_status in ["账期未到", "账期已结"] else None

            else:
                # 旧版/v1/v2：使用"初始状态"和"先采后付"字段
                payment_method = "先采后付" if row.get("先采后付") == "先采后付" else "已付款"
                order_status = str(row["初始状态"]).strip() if pd.notna(row["初始状态"]) else "已付款"

                # 映射到新的payment_status
                if payment_method == "先采后付":
                    payment_status = "账期未到"  # 默认为账期未到
                else:
                    payment_status = "即时付款"

                receive_date = None

            order_data = {
                "order_no": order_no,
                "product_name": str(row["产品名称"]).strip(),
                "spec": str(row["规格"]).strip() if pd.notna(row["规格"]) else None,
                "purchase_amount": float(row["采购金额"]),
                "supplier": str(row["供应商"]).strip() if pd.notna(row["供应商"]) else None,
                "order_date": order_date,
                "order_status": order_status,  # DEPRECATED
                "payment_method": payment_method,  # DEPRECATED
                "payment_status": payment_status,  # 新字段
                "receive_date": receive_date,  # 确认收货日期
                "record_time": record_time,
            }

            orders.append(order_data)

        except Exception as e:
            errors.append({"row": idx, "error": str(e)})

    return orders, errors


def export_orders_to_excel(orders):
    """导出订单到Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "订单列表"

    # 表头
    headers = ["订单编号", "产品名称", "规格", "采购金额", "供应商", "订单日期", "订单状态", "支付方式", "记录时间"]
    ws.append(headers)

    # 设置表头样式
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 填充数据
    for order in orders:
        ws.append([
            order.order_no,
            order.product_name,
            order.spec if order.spec else "",
            float(order.purchase_amount),
            order.supplier if order.supplier else "",
            order.order_date.strftime("%Y-%m-%d"),
            order.order_status,
            order.payment_method,
            order.record_time.strftime("%Y-%m-%d %H:%M:%S"),
        ])

    # 调整列宽
    ws.column_dimensions["A"].width = 25  # 订单编号
    ws.column_dimensions["B"].width = 30  # 产品名称
    ws.column_dimensions["C"].width = 20  # 规格
    ws.column_dimensions["D"].width = 15  # 采购金额
    ws.column_dimensions["E"].width = 30  # 供应商
    ws.column_dimensions["F"].width = 15  # 订单日期
    ws.column_dimensions["G"].width = 12  # 订单状态
    ws.column_dimensions["H"].width = 12  # 支付方式
    ws.column_dimensions["I"].width = 20  # 记录时间

    # 保存到BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output.getvalue()
